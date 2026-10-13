# -*- coding: utf-8 -*-
import io, re, os, sys
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

BASE = sys.argv[1]
SRC = ['README.md','01_시스템_구조.md','02_판정_엔진.md','03_기관인증과_청구서.md',
       '04_블록체인.md','05_백엔드와_실사용모드.md','06_확장_기능.md','07_데모_시나리오_17종.md']

doc = Document()
# 한글 폰트 고정
st = doc.styles['Normal']
st.font.name = '맑은 고딕'; st.font.size = Pt(9.5)
st.element.rPr.rFonts.set(qn('w:eastAsia'), '맑은 고딕')
st.paragraph_format.space_after = Pt(4)
st.paragraph_format.line_spacing = 1.3
for s in doc.sections:
    s.page_width, s.page_height = Cm(21), Cm(29.7)
    s.left_margin = s.right_margin = Cm(1.8)
    s.top_margin = s.bottom_margin = Cm(1.8)

def runs(p, text, base_bold=False):
    """**bold** / `code` 를 run 으로 분리"""
    for tok in re.split(r'(\*\*[^*]+\*\*|`[^`]+`)', text):
        if not tok: continue
        if tok.startswith('**') and tok.endswith('**'):
            r = p.add_run(tok[2:-2]); r.bold = True
        elif tok.startswith('`') and tok.endswith('`'):
            r = p.add_run(tok[1:-1]); r.font.name = 'Consolas'
            r.element.rPr.rFonts.set(qn('w:eastAsia'), 'Consolas')
            r.font.size = Pt(8.8)
        else:
            r = p.add_run(re.sub(r'\[([^\]]+)\]\(([^)]+)\)', r'\1 (\2)', tok))
        if base_bold: r.bold = True

def para(text, style=None, size=None, space_before=None):
    p = doc.add_paragraph(style=style)
    if space_before: p.paragraph_format.space_before = Pt(space_before)
    runs(p, text)
    if size:
        for r in p.runs: r.font.size = Pt(size)
    return p

def heading(text, lv):
    sizes = {1: 17, 2: 13, 3: 11, 4: 10}
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0 if lv == 1 else (14 if lv == 2 else 10))
    p.paragraph_format.space_after = Pt(5 if lv == 1 else 4)
    runs(p, text, base_bold=True)
    for r in p.runs:
        r.font.size = Pt(sizes.get(lv, 10)); r.bold = True
        if lv <= 2: r.font.color.rgb = RGBColor(0x11, 0x11, 0x11)
    return p

def cells_of(s): return [c.strip() for c in s.strip().strip('|').split('|')]

for idx, f in enumerate(SRC):
    md = io.open(os.path.join(BASE, f), encoding='utf-8').read().replace('\r\n', '\n')
    if idx > 0:
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    lines = md.split('\n'); i, n = 0, len(lines)
    while i < n:
        ln = lines[i]
        if ln.startswith('```'):
            i += 1; buf = []
            while i < n and not lines[i].startswith('```'):
                buf.append(lines[i]); i += 1
            i += 1
            p = doc.add_paragraph(); p.paragraph_format.left_indent = Cm(0.4)
            p.paragraph_format.space_before = Pt(3); p.paragraph_format.line_spacing = 1.15
            r = p.add_run('\n'.join(buf)); r.font.name = 'Consolas'
            r.element.rPr.rFonts.set(qn('w:eastAsia'), 'Consolas'); r.font.size = Pt(8.3)
            continue
        if ln.strip().startswith('|') and i + 1 < n and re.match(r'^\s*\|[\s:|-]+\|\s*$', lines[i+1]):
            head = cells_of(ln); i += 2; rows = []
            while i < n and lines[i].strip().startswith('|'):
                rows.append(cells_of(lines[i])); i += 1
            cols = max(len(head), max((len(r) for r in rows), default=0))
            t = doc.add_table(rows=1, cols=cols); t.style = 'Table Grid'
            t.alignment = WD_TABLE_ALIGNMENT.LEFT
            for k in range(cols):
                c = t.rows[0].cells[k]; c.paragraphs[0].text = ''
                runs(c.paragraphs[0], head[k] if k < len(head) else '', base_bold=True)
                for r in c.paragraphs[0].runs: r.font.size = Pt(8.8)
            for row in rows:
                rc = t.add_row().cells
                for k in range(cols):
                    rc[k].paragraphs[0].text = ''
                    runs(rc[k].paragraphs[0], row[k] if k < len(row) else '')
                    for r in rc[k].paragraphs[0].runs: r.font.size = Pt(8.8)
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
            continue
        m = re.match(r'^(#{1,4})\s+(.*)$', ln)
        if m:
            heading(m.group(2), len(m.group(1))); i += 1; continue
        if re.match(r'^---+\s*$', ln): i += 1; continue
        if ln.startswith('>'):
            buf = []
            while i < n and lines[i].startswith('>'):
                buf.append(lines[i].lstrip('>').strip()); i += 1
            p = para(' '.join(buf), size=9)
            p.paragraph_format.left_indent = Cm(0.5)
            for r in p.runs: r.italic = True
            continue
        if re.match(r'^\s*[-*]\s+', ln):
            while i < n and re.match(r'^\s*[-*]\s+', lines[i]):
                para(re.sub(r'^\s*[-*]\s+', '', lines[i]), style='List Bullet'); i += 1
            continue
        if re.match(r'^\s*\d+\.\s+', ln):
            while i < n and re.match(r'^\s*\d+\.\s+', lines[i]):
                para(re.sub(r'^\s*\d+\.\s+', '', lines[i]), style='List Number'); i += 1
            continue
        if ln.strip():
            buf = []
            while i < n and lines[i].strip() and not lines[i].startswith(('#', '>', '|', '```')) \
                  and not re.match(r'^\s*[-*]\s+|^\s*\d+\.\s+|^---+\s*$', lines[i]):
                buf.append(lines[i].strip()); i += 1
            para(' '.join(buf)); continue
        i += 1

out = os.path.join(BASE, '배포본', 'CarePass_기능설명서.docx')
doc.save(out)
print('DOCX 저장:', out)
